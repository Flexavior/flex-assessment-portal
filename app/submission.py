"""Submission document extraction — rich text from student uploads."""
from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass, field

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

from app.config import GRADING_CHUNK_OVERLAP, MAX_CHARS_PER_ANSWER, MAX_SUBMISSION_CHARS


class SubmissionTooLargeError(ValueError):
    """Extracted submission text exceeds configured grading limits."""

    def __init__(self, char_count: int, max_chars: int):
        self.char_count = char_count
        self.max_chars = max_chars
        super().__init__(
            f"Submission text is {char_count:,} characters — exceeds the maximum "
            f"supported limit of {max_chars:,}. Shorten the file or raise "
            "MAX_SUBMISSION_CHARS in .env."
        )

TEXT_EXTENSIONS = {".txt", ".md"}
DOCX_EXTENSIONS = {".docx", ".doc"}
PDF_EXTENSIONS = {".pdf"}
PPT_EXTENSIONS = {".pptx", ".ppt"}

_HEADING_RE = re.compile(
    r"^(?:"
    r"(?:\d+\.)+\d*\s+\S|"           # 3.1 Title
    r"Section\s+\d+|"
    r"Q\d+[:.]?\s+|"
    r"[A-Z][A-Za-z0-9\s/&\-]{4,60}$"  # Title Case heading
    r")",
    re.MULTILINE,
)


@dataclass
class SubmissionContent:
    text: str
    word_count: int = 0
    char_count: int = 0
    section_headings: list[str] = field(default_factory=list)
    embedded_images: int = 0
    source_file: str = ""
    extraction_notes: str = ""

    def summary_for_prompt(self) -> str:
        parts = [
            f"Extracted {self.word_count} words ({self.char_count:,} characters) from {self.source_file}.",
        ]
        if self.section_headings:
            parts.append(
                "Detected section headings: "
                + "; ".join(self.section_headings[:12])
            )
        if self.embedded_images:
            parts.append(
                f"Embedded images/diagrams: {self.embedded_images} "
                "(evaluate diagram topics from surrounding text and captions)."
            )
        if self.extraction_notes:
            parts.append(self.extraction_notes)
        return " ".join(parts)

    def validate_grading_size(self) -> None:
        if self.char_count > MAX_SUBMISSION_CHARS:
            raise SubmissionTooLargeError(self.char_count, MAX_SUBMISSION_CHARS)

    def grading_chunks(self) -> list[str]:
        """Split extracted text into model-sized chunks. Never silently truncates."""
        self.validate_grading_size()
        return split_for_grading(self.text, MAX_CHARS_PER_ANSWER, GRADING_CHUNK_OVERLAP)


def split_for_grading(text: str, max_chars: int, overlap: int = 0) -> list[str]:
    if not text:
        return [""]
    if len(text) <= max_chars:
        return [text]

    sections = _split_by_headings(text)
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    def flush():
        nonlocal current, current_len
        if current:
            chunks.append("\n\n".join(current))
            current = []
            current_len = 0

    def append_hard_parts(section: str):
        step = max(1, max_chars - max(overlap, 0))
        start = 0
        while start < len(section):
            end = min(len(section), start + max_chars)
            chunks.append(section[start:end])
            if end >= len(section):
                break
            start = end - overlap if overlap else end

    for section in sections:
        section = section.strip()
        if not section:
            continue
        extra = len(section) + (2 if current else 0)
        if len(section) > max_chars:
            flush()
            append_hard_parts(section)
            continue
        if current and current_len + extra > max_chars:
            flush()
        current.append(section)
        current_len += extra

    flush()
    return chunks or [text[:max_chars]]


def _split_by_headings(text: str) -> list[str]:
    lines = text.splitlines()
    if not lines:
        return [text]

    sections: list[str] = []
    current: list[str] = []
    for line in lines:
        stripped = line.strip()
        is_heading = bool(
            stripped
            and len(stripped) <= 120
            and (_HEADING_RE.match(stripped) or (stripped[0].isupper() and len(stripped.split()) <= 10 and not stripped.endswith(".")))
        )
        if is_heading and current:
            sections.append("\n".join(current))
            current = [line]
        else:
            current.append(line)
    if current:
        sections.append("\n".join(current))
    return sections if sections else [text]


def _sanitize(text: str) -> str:
    if not text:
        return ""
    cleaned = []
    for ch in text:
        if ch in "\n\r\t":
            cleaned.append(ch)
        elif ch.isprintable():
            cleaned.append(ch)
        else:
            cleaned.append(" ")
    text = "".join(cleaned)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _word_count(text: str) -> int:
    return len(re.findall(r"\b\w+\b", text or ""))


def _detect_headings(text: str) -> list[str]:
    headings = []
    for line in text.splitlines():
        line = line.strip()
        if not line or len(line) > 120:
            continue
        if _HEADING_RE.match(line) or (
            line[0].isupper() and len(line.split()) <= 10 and not line.endswith(".")
        ):
            if line not in headings:
                headings.append(line)
    return headings[:20]


def _count_docx_images(path: str) -> int:
    try:
        with zipfile.ZipFile(path) as zf:
            return sum(
                1 for name in zf.namelist()
                if name.startswith("word/media/")
            )
    except Exception:
        return 0


def read_docx(path: str) -> str:
    doc = Document(path)
    parts = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    body = doc.element.body
    for child in body.iter():
        if child.tag == qn("w:txbxContent"):
            for p in child.iter(qn("w:p")):
                texts = [node.text for node in p.iter(qn("w:t")) if node.text]
                if texts:
                    parts.append("".join(texts))

    return _sanitize("\n\n".join(parts))


def read_pdf(path: str) -> str:
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                chunks.append(text)
    return _sanitize("\n".join(chunks))


def read_ppt(path: str) -> str:
    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                chunks.append(shape.text)
    return _sanitize("\n".join(chunks))


def read_txt(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as handle:
        return _sanitize(handle.read())


def extract_submission(path: str) -> SubmissionContent:
    """Extract gradable text. Never use LlamaIndex SimpleDirectoryReader for students."""
    path = os.path.abspath(path)
    name = os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    images = 0
    notes = ""

    if ext in DOCX_EXTENSIONS:
        text = read_docx(path)
        images = _count_docx_images(path)
        method = "docx-structured"
    elif ext in PDF_EXTENSIONS:
        text = read_pdf(path)
        method = "pdf-text"
    elif ext in PPT_EXTENSIONS:
        text = read_ppt(path)
        method = "pptx-text"
    elif ext in TEXT_EXTENSIONS:
        text = read_txt(path)
        method = "plain-text"
    else:
        text = ""
        method = "unsupported"
        notes = f"Unsupported extension: {ext}"

    words = _word_count(text)
    if words < 50 and images > 0:
        notes = (
            "Low text but images present — score diagram topics if captions or "
            "surrounding prose describe the figure."
        )
    elif words < 50:
        notes = "Very little text extracted — verify the file is not scanned-only PDF."

    return SubmissionContent(
        text=text,
        word_count=words,
        char_count=len(text),
        section_headings=_detect_headings(text),
        embedded_images=images,
        source_file=name,
        extraction_notes=notes,
    )
